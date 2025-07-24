"""
Document to PDF Conversion Utilities

This module provides utilities for converting various document formats to PDF.
It tries multiple conversion methods in order of quality/reliability.
"""

import logging
import os
import subprocess
import tempfile
from pathlib import Path
from typing import Optional, Tuple
import platform

logger = logging.getLogger(__name__)

# Check available conversion libraries
UNOCONV_AVAILABLE = False
DOCX2PDF_AVAILABLE = False
PYPANDOC_AVAILABLE = False
PYTHON_DOCX_AVAILABLE = False
OPENPYXL_AVAILABLE = False
REPORTLAB_AVAILABLE = False

# Check if unoconv is available (requires LibreOffice)
try:
    result = subprocess.run(['unoconv', '--version'], capture_output=True, text=True)
    if result.returncode == 0:
        UNOCONV_AVAILABLE = True
        logger.info("✅ unoconv is available for document conversion")
except (FileNotFoundError, subprocess.SubprocessError):
    logger.info("❌ unoconv not available (install LibreOffice for best conversion quality)")

# Check Python libraries
try:
    import docx2pdf
    DOCX2PDF_AVAILABLE = True
    logger.info("✅ docx2pdf is available")
except ImportError:
    logger.info("❌ docx2pdf not available")

try:
    import pypandoc
    PYPANDOC_AVAILABLE = True
    logger.info("✅ pypandoc is available")
except ImportError:
    logger.info("❌ pypandoc not available")

try:
    import docx
    PYTHON_DOCX_AVAILABLE = True
    logger.info("✅ python-docx is available")
except ImportError:
    logger.info("❌ python-docx not available")

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
    logger.info("✅ openpyxl is available")
except ImportError:
    logger.info("❌ openpyxl not available")

try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    REPORTLAB_AVAILABLE = True
    logger.info("✅ reportlab is available")
except ImportError:
    logger.info("❌ reportlab not available")


class DocumentConverter:
    """Handles conversion of various document formats to PDF"""
    
    @staticmethod
    def convert_to_pdf(input_path: str, output_path: Optional[str] = None) -> Tuple[bool, Optional[str]]:
        """
        Convert a document to PDF
        
        Args:
            input_path: Path to the input document
            output_path: Optional path for the output PDF. If not provided, generates one.
            
        Returns:
            Tuple of (success: bool, pdf_path: str or None)
        """
        input_path = Path(input_path)
        
        if not input_path.exists():
            logger.error(f"Input file does not exist: {input_path}")
            return False, None
            
        # Generate output path if not provided
        if output_path is None:
            output_path = input_path.with_suffix('.pdf')
        else:
            output_path = Path(output_path)
            
        # Get file extension
        file_extension = input_path.suffix.lower()
        
        # If already PDF, just return the input path
        if file_extension == '.pdf':
            logger.info(f"File is already PDF: {input_path}")
            return True, str(input_path)
            
        # Try conversion based on file type
        logger.info(f"Converting {file_extension} to PDF: {input_path}")
        
        # For Word documents
        if file_extension in ['.doc', '.docx']:
            return DocumentConverter._convert_word_to_pdf(input_path, output_path)
            
        # For Excel documents
        elif file_extension in ['.xls', '.xlsx']:
            return DocumentConverter._convert_excel_to_pdf(input_path, output_path)
            
        # For PowerPoint documents
        elif file_extension in ['.ppt', '.pptx']:
            return DocumentConverter._convert_powerpoint_to_pdf(input_path, output_path)
            
        # For text files
        elif file_extension in ['.txt', '.md', '.rst']:
            return DocumentConverter._convert_text_to_pdf(input_path, output_path)
            
        # For HTML files
        elif file_extension in ['.html', '.htm']:
            return DocumentConverter._convert_html_to_pdf(input_path, output_path)
            
        else:
            logger.warning(f"Unsupported file type for PDF conversion: {file_extension}")
            return False, None
    
    @staticmethod
    def _convert_word_to_pdf(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Convert Word document to PDF"""
        
        # Method 1: Try unoconv (best quality)
        if UNOCONV_AVAILABLE:
            try:
                logger.info("Converting Word to PDF using unoconv...")
                result = subprocess.run(
                    ['unoconv', '-f', 'pdf', '-o', str(output_path), str(input_path)],
                    capture_output=True,
                    text=True,
                    timeout=60
                )
                if result.returncode == 0 and output_path.exists():
                    logger.info(f"✅ Successfully converted Word to PDF: {output_path}")
                    return True, str(output_path)
                else:
                    logger.error(f"unoconv failed: {result.stderr}")
            except Exception as e:
                logger.error(f"unoconv conversion error: {e}")
        
        # Method 2: Try docx2pdf (Windows/macOS only)
        if DOCX2PDF_AVAILABLE and platform.system() in ['Windows', 'Darwin']:
            try:
                logger.info("Converting Word to PDF using docx2pdf...")
                import docx2pdf
                docx2pdf.convert(str(input_path), str(output_path))
                if output_path.exists():
                    logger.info(f"✅ Successfully converted Word to PDF: {output_path}")
                    return True, str(output_path)
            except Exception as e:
                logger.error(f"docx2pdf conversion error: {e}")
        
        # Method 3: Try pypandoc
        if PYPANDOC_AVAILABLE:
            try:
                logger.info("Converting Word to PDF using pypandoc...")
                import pypandoc
                pypandoc.convert_file(str(input_path), 'pdf', outputfile=str(output_path))
                if output_path.exists():
                    logger.info(f"✅ Successfully converted Word to PDF: {output_path}")
                    return True, str(output_path)
            except Exception as e:
                logger.error(f"pypandoc conversion error: {e}")
        
        # Method 4: Fallback - extract text and create simple PDF
        if PYTHON_DOCX_AVAILABLE and REPORTLAB_AVAILABLE:
            try:
                logger.info("Converting Word to PDF using python-docx + reportlab...")
                return DocumentConverter._word_to_pdf_fallback(input_path, output_path)
            except Exception as e:
                logger.error(f"Fallback Word to PDF conversion error: {e}")
        
        logger.error("❌ No method available to convert Word to PDF")
        return False, None
    
    @staticmethod
    def _word_to_pdf_fallback(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Fallback method to convert Word to PDF by extracting text"""
        import docx
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import inch
        
        # Read Word document
        doc = docx.Document(str(input_path))
        
        # Create PDF
        pdf_doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=18
        )
        
        # Container for the 'Flowable' objects
        elements = []
        styles = getSampleStyleSheet()
        
        # Process paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                # Determine style based on paragraph style
                if para.style.name.startswith('Heading'):
                    style = styles['Heading1']
                elif para.style.name == 'Title':
                    style = styles['Title']
                else:
                    style = styles['Normal']
                
                # Create paragraph
                p = Paragraph(para.text, style)
                elements.append(p)
                elements.append(Spacer(1, 12))
        
        # Build PDF
        pdf_doc.build(elements)
        
        if output_path.exists():
            logger.info(f"✅ Successfully converted Word to PDF (fallback): {output_path}")
            return True, str(output_path)
        
        return False, None
    
    @staticmethod
    def _convert_excel_to_pdf(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Convert Excel document to PDF"""
        
        # Method 1: Try unoconv (best quality)
        if UNOCONV_AVAILABLE:
            try:
                logger.info("Converting Excel to PDF using unoconv...")
                result = subprocess.run(
                    ['unoconv', '-f', 'pdf', '-o', str(output_path), str(input_path)],
                    capture_output=True,
                    text=True,
                    timeout=60
                )
                if result.returncode == 0 and output_path.exists():
                    logger.info(f"✅ Successfully converted Excel to PDF: {output_path}")
                    return True, str(output_path)
            except Exception as e:
                logger.error(f"unoconv Excel conversion error: {e}")
        
        # Method 2: Fallback - extract data and create simple PDF
        if OPENPYXL_AVAILABLE and REPORTLAB_AVAILABLE:
            try:
                logger.info("Converting Excel to PDF using openpyxl + reportlab...")
                return DocumentConverter._excel_to_pdf_fallback(input_path, output_path)
            except Exception as e:
                logger.error(f"Fallback Excel to PDF conversion error: {e}")
        
        logger.error("❌ No method available to convert Excel to PDF")
        return False, None
    
    @staticmethod
    def _excel_to_pdf_fallback(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Fallback method to convert Excel to PDF"""
        import openpyxl
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        
        # Read Excel file
        wb = openpyxl.load_workbook(str(input_path), read_only=True)
        
        # Create PDF
        pdf_doc = SimpleDocTemplate(
            str(output_path),
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=18
        )
        
        elements = []
        styles = getSampleStyleSheet()
        
        # Process each sheet
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # Add sheet title
            elements.append(Paragraph(f"Sheet: {sheet_name}", styles['Heading1']))
            elements.append(Spacer(1, 12))
            
            # Extract data from sheet
            data = []
            for row in ws.iter_rows(values_only=True):
                # Filter out completely empty rows
                if any(cell is not None for cell in row):
                    # Convert None to empty string
                    data.append([str(cell) if cell is not None else '' for cell in row])
            
            if data:
                # Create table
                table = Table(data)
                
                # Add style to table
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                
                elements.append(table)
                elements.append(Spacer(1, 20))
        
        # Build PDF
        pdf_doc.build(elements)
        
        if output_path.exists():
            logger.info(f"✅ Successfully converted Excel to PDF (fallback): {output_path}")
            return True, str(output_path)
        
        return False, None
    
    @staticmethod
    def _convert_powerpoint_to_pdf(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Convert PowerPoint to PDF"""
        
        # Method 1: Try unoconv (best quality)
        if UNOCONV_AVAILABLE:
            try:
                logger.info("Converting PowerPoint to PDF using unoconv...")
                result = subprocess.run(
                    ['unoconv', '-f', 'pdf', '-o', str(output_path), str(input_path)],
                    capture_output=True,
                    text=True,
                    timeout=60
                )
                if result.returncode == 0 and output_path.exists():
                    logger.info(f"✅ Successfully converted PowerPoint to PDF: {output_path}")
                    return True, str(output_path)
            except Exception as e:
                logger.error(f"unoconv PowerPoint conversion error: {e}")
        
        logger.error("❌ No method available to convert PowerPoint to PDF")
        return False, None
    
    @staticmethod
    def _convert_text_to_pdf(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Convert text file to PDF"""
        
        if REPORTLAB_AVAILABLE:
            try:
                logger.info("Converting text to PDF using reportlab...")
                from reportlab.pdfgen import canvas
                from reportlab.lib.pagesizes import letter
                
                # Read text file
                with open(input_path, 'r', encoding='utf-8') as f:
                    text = f.read()
                
                # Create PDF
                c = canvas.Canvas(str(output_path), pagesize=letter)
                width, height = letter
                
                # Set font
                c.setFont("Courier", 10)
                
                # Write text
                y = height - 40
                for line in text.split('\n'):
                    if y < 40:  # New page if needed
                        c.showPage()
                        c.setFont("Courier", 10)
                        y = height - 40
                    
                    # Truncate long lines
                    if len(line) > 80:
                        line = line[:80] + '...'
                    
                    c.drawString(40, y, line)
                    y -= 12
                
                c.save()
                
                if output_path.exists():
                    logger.info(f"✅ Successfully converted text to PDF: {output_path}")
                    return True, str(output_path)
            except Exception as e:
                logger.error(f"Text to PDF conversion error: {e}")
        
        logger.error("❌ No method available to convert text to PDF")
        return False, None
    
    @staticmethod
    def _convert_html_to_pdf(input_path: Path, output_path: Path) -> Tuple[bool, Optional[str]]:
        """Convert HTML to PDF"""
        
        # Try pypandoc if available
        if PYPANDOC_AVAILABLE:
            try:
                logger.info("Converting HTML to PDF using pypandoc...")
                import pypandoc
                pypandoc.convert_file(str(input_path), 'pdf', outputfile=str(output_path))
                if output_path.exists():
                    logger.info(f"✅ Successfully converted HTML to PDF: {output_path}")
                    return True, str(output_path)
            except Exception as e:
                logger.error(f"pypandoc HTML conversion error: {e}")
        
        logger.error("❌ No method available to convert HTML to PDF")
        return False, None


# Module-level function for easy access
def convert_document_to_pdf(input_path: str, output_path: Optional[str] = None) -> Tuple[bool, Optional[str]]:
    """
    Convert a document to PDF
    
    Args:
        input_path: Path to the input document
        output_path: Optional path for the output PDF
        
    Returns:
        Tuple of (success: bool, pdf_path: str or None)
    """
    return DocumentConverter.convert_to_pdf(input_path, output_path)